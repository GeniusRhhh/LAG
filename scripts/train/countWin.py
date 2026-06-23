import pandas as pd
import numpy as np
import matplotlib.pyplot as plt

# 设置中文字体支持
plt.rcParams['font.sans-serif'] = ['SimHei']  # 指定默认字体为黑体
plt.rcParams['axes.unicode_minus'] = False  # 解决保存图像时负号'-'显示为方块的问题


def filter_useful_rows(df):
    """
    过滤出有用的行（每个轮次的第一行）

    参数:
        df: 原始DataFrame

    返回:
        filtered_df: 过滤后只包含有用行的DataFrame
    """
    # 创建一个新列标记第一行和第二行
    df['row_type'] = df.index % 2  # 偶数索引为0（第一行），奇数索引为1（第二行）

    # 只保留第一行（row_type=0）
    filtered_df = df[df['row_type'] == 0].copy()

    # 删除辅助列
    filtered_df.drop('row_type', axis=1, inplace=True)

    return filtered_df


def process_combat_data(df):
    """
    处理对战数据，按回合范围分组并计算胜率统计

    参数:
        df: 过滤后只包含有用行的DataFrame

    返回:
        summary_df: 按回合范围汇总的统计数据
    """
    # 确保数据类型正确
    df['模型轮次'] = df['模型轮次'].astype(int)

    # 定义回合范围
    ranges = [
        (0, 200),
        (200, 400),
        (400, 600),
        (600, 714)  # 调整为您的最大轮次
    ]

    # 创建结果列表
    results = []

    for start, end in ranges:
        # 筛选该范围内的数据
        range_df = df[(df['模型轮次'] >= start) & (df['模型轮次'] < end)]

        if len(range_df) > 0:
            # 计算胜率、败率和平局率
            total_games = len(range_df)
            wins = sum(range_df['胜负结果'] == '胜利')
            losses = sum(range_df['胜负结果'] == '失败')
            draws = sum(range_df['胜负结果'] == '平局')

            win_rate = (wins / total_games) * 100
            lose_rate = (losses / total_games) * 100
            draw_rate = (draws / total_games) * 100

            # 计算平均奖励差
            avg_reward_diff = range_df['奖励差'].mean()

            # 添加到结果列表
            results.append({
                '回合范围': f"{start}-{end}",
                '对战场次': total_games,
                '红方胜率': win_rate,
                '红方败率': lose_rate,
                '平局率': draw_rate,
                '平均奖励差': avg_reward_diff
            })

    # 创建汇总DataFrame
    summary_df = pd.DataFrame(results)

    # 格式化百分比
    summary_df['红方胜率'] = summary_df['红方胜率'].apply(lambda x: f"{x:.1f}%")
    summary_df['红方败率'] = summary_df['红方败率'].apply(lambda x: f"{x:.1f}%")
    summary_df['平局率'] = summary_df['平局率'].apply(lambda x: f"{x:.1f}%")

    return summary_df


def create_visualization(summary_df):
    """
    基于汇总数据创建可视化图表

    参数:
        summary_df: 按回合范围汇总的统计数据
    """
    # 将百分比字符串转换回数字
    summary_df['红方胜率_num'] = summary_df['红方胜率'].str.rstrip('%').astype(float)
    summary_df['红方败率_num'] = summary_df['红方败率'].str.rstrip('%').astype(float)
    summary_df['平局率_num'] = summary_df['平局率'].str.rstrip('%').astype(float)

    # 创建柱状图
    plt.figure(figsize=(12, 8))

    x = range(len(summary_df))
    width = 0.25

    plt.bar([i - width for i in x], summary_df['红方胜率_num'], width=width, label='红方胜率', color='red')
    plt.bar(x, summary_df['红方败率_num'], width=width, label='红方败率', color='blue')
    plt.bar([i + width for i in x], summary_df['平局率_num'], width=width, label='平局率', color='gray')

    plt.xlabel('训练回合范围')
    plt.ylabel('百分比')
    plt.title('PPO 算法与 Pursue 基线模型的胜率统计')
    plt.xticks(x, [f"{r} 回合" for r in summary_df['回合范围']])
    plt.legend()

    plt.tight_layout()
    plt.savefig('ppo_vs_pursue_statistics.png', dpi=300)
    plt.show()


def save_results_to_excel(summary_df, file_path='PPO_vs_Pursue_Results.xlsx'):
    """
    将汇总结果保存到Excel文件

    参数:
        summary_df: 汇总统计数据
        file_path: 保存的Excel文件路径
    """
    try:
        # 直接保存为Excel
        summary_df.to_excel(file_path, sheet_name='胜率统计', index=False)
        print(f"结果已保存至 {file_path}")

        # 如果需要更美观的格式，可以添加以下代码
        # 需要安装openpyxl: pip install openpyxl
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment

            # 打开文件进行美化
            wb = openpyxl.load_workbook(file_path)
            ws = wb['胜率统计']

            # 插入标题行
            ws.insert_rows(1)
            title_cell = ws.cell(row=1, column=1)
            title_cell.value = "表 3.1 PPO 算法与 Pursue 基线模型的胜率统计"
            title_cell.font = Font(bold=True, size=14)

            # 合并标题单元格
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(summary_df.columns))

            # 设置列宽
            for i, col in enumerate(ws.columns, 1):
                if i == 1:  # 回合范围列
                    ws.column_dimensions[col[0].column_letter].width = 15
                else:
                    ws.column_dimensions[col[0].column_letter].width = 12

            # 保存美化后的文件
            wb.save(file_path)
            print(f"文件已美化并保存")
        except:
            print("无法进行Excel美化，但文件已基本保存")

    except Exception as e:
        print(f"保存Excel文件时出错: {e}")
        # 尝试保存为CSV作为备选
        try:
            summary_df.to_csv(file_path.replace('.xlsx', '.csv'), index=False, encoding='utf-8-sig')
            print(f"已保存为CSV: {file_path.replace('.xlsx', '.csv')}")
        except:
            print("无法保存文件")


def process_excel_file(file_path, sheet_name=0):
    """
    处理Excel文件中的对战数据

    参数:
        file_path: Excel文件路径
        sheet_name: 工作表名称或索引
    """
    try:
        # 读取Excel文件
        print(f"正在读取文件: {file_path}")
        df = pd.read_excel(file_path, sheet_name=sheet_name)

        # 显示原始数据的前几行
        print("原始数据预览:")
        print(df.head())

        # 过滤有用的行
        print("正在过滤有用行...")
        filtered_df = filter_useful_rows(df)
        print(f"过滤后数据行数: {len(filtered_df)}")

        # 显示过滤后的前几行
        print("过滤后数据预览:")
        print(filtered_df.head())

        # 处理数据
        print("正在处理数据...")
        summary_df = process_combat_data(filtered_df)

        # 显示处理结果
        print("\n处理结果:")
        print(summary_df.to_string(index=False))

        # 创建可视化
        print("正在创建可视化...")
        create_visualization(summary_df)

        # 保存结果
        save_results_to_excel(summary_df)

        print("处理完成!")
        return summary_df

    except Exception as e:
        print(f"处理文件时出错: {e}")
        import traceback
        traceback.print_exc()
        return None


# 如果数据已经是按特定模式排列（如每隔一行是有用数据），也可以使用以下方法
def filter_by_pattern(df):
    """
    根据特定模式过滤数据（每隔一行取数据）

    参数:
        df: 原始DataFrame

    返回:
        filtered_df: 过滤后只包含有用行的DataFrame
    """
    # 只保留奇数索引行或偶数索引行
    return df.iloc[::2].copy()  # 每隔一行取一次，从第0行开始


# 或者通过识别内容特征来过滤
def filter_by_content(df):
    """
    通过内容特征过滤数据（不包含'汇总'字样的行）

    参数:
        df: 原始DataFrame

    返回:
        filtered_df: 过滤后只包含有用行的DataFrame
    """
    # 假设第二列中包含'汇总'字样的行需要被过滤掉
    # 这取决于您数据的具体特征
    second_col = df.columns[1]  # 获取第二列名称
    return df[~df[second_col].astype(str).str.contains('汇总')].copy()


# 主函数
if __name__ == "__main__":
    # 文件路径
    file_path = "ppo_vs_pursue_test.xlsx"

    # 处理文件
    process_excel_file(file_path)